#!/usr/bin/env python3
"""XLS → XLSX 変換ツール（GUIワンボタン版）"""

import tkinter as tk
from tkinter import filedialog, messagebox
import os
import xlrd
import openpyxl


def convert(src: str, dst: str | None = None) -> str:
    if dst is None:
        dst = os.path.splitext(src)[0] + ".xlsx"

    wb_xls = xlrd.open_workbook(src)
    wb_xlsx = openpyxl.Workbook()
    wb_xlsx.remove(wb_xlsx.active)

    for sheet_name in wb_xls.sheet_names():
        ws_xls = wb_xls.sheet_by_name(sheet_name)
        ws_xlsx = wb_xlsx.create_sheet(title=sheet_name)

        for row_idx in range(ws_xls.nrows):
            for col_idx in range(ws_xls.ncols):
                cell = ws_xls.cell(row_idx, col_idx)
                value = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        value = xlrd.xldate_as_datetime(value, wb_xls.datemode)
                    except Exception:
                        pass
                ws_xlsx.cell(row=row_idx + 1, column=col_idx + 1, value=value)

    wb_xlsx.save(dst)
    return dst


def on_click():
    path = filedialog.askopenfilename(
        title="変換するXLSファイルを選択",
        filetypes=[("Excel 97-2003", "*.xls")],
    )
    if not path:
        return

    try:
        output = convert(path)
        messagebox.showinfo("完了", f"変換しました！\n\n{output}")
    except Exception as e:
        messagebox.showerror("エラー", str(e))


root = tk.Tk()
root.title("XLS → XLSX 変換ツール")
root.resizable(False, False)

frame = tk.Frame(root, padx=40, pady=30)
frame.pack()

tk.Label(frame, text="XLSファイルをXLSXに変換します", font=("", 12)).pack(pady=(0, 20))

btn = tk.Button(
    frame,
    text="ファイルを選んで変換",
    font=("", 13, "bold"),
    bg="#4CAF50",
    fg="white",
    padx=20,
    pady=10,
    cursor="hand2",
    command=on_click,
)
btn.pack()

tk.Label(frame, text="変換後のファイルは元ファイルと同じフォルダに保存されます", font=("", 9), fg="gray").pack(pady=(15, 0))

root.mainloop()
